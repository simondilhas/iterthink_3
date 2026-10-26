"""An edited delivery file comes back as one proposal to the open source."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

from iterthink.services.bridge_real import RealBridge


def _bridge(tmp_path: Path, project_dir: Path) -> tuple[RealBridge, str]:
    bridge = RealBridge(state_dir=tmp_path / "state", machine="m-test")
    return bridge, bridge.make_project(project_dir).id


def test_returned_docx_becomes_proposal_to_markdown(tmp_path: Path, project_dir: Path) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    returned = project_dir / "from-client.docx"
    document = Document()
    document.add_paragraph("Die Türen im EG sind EI90.")
    document.add_paragraph("Fluchtwege 1.2 m.")
    document.save(str(returned))
    original = returned.read_bytes()
    info = bridge.import_returned(project, returned, "Pflichtenheft.md", "user:ada")
    assert returned.read_bytes() == original
    assert info.paths == ("Pflichtenheft.md",)
    assert info.reason == "changes from from-client.docx"
    assert info.status == "open"
    diffs = bridge.proposal_diff(project, info.id)
    assert len(diffs) == 1
    assert diffs[0].path == "Pflichtenheft.md"
    after = "\n".join(block.after for block in diffs[0].blocks)
    assert "EI90" in after
    assert "EI30" not in after


def test_returned_xlsx_becomes_proposal_to_csv(tmp_path: Path, project_dir: Path) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    returned = project_dir / "back.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["id", "geschoss", "typ", "brandschutz"])
    sheet.append(["T01", "EG", "Holz", "EI90"])
    sheet.append(["T02", "EG", "Glas", ""])
    book.save(returned)
    original = returned.read_bytes()
    info = bridge.import_returned(project, returned, "Türliste.csv", "user:ada")
    assert returned.read_bytes() == original
    assert info.paths == ("Türliste.csv",)
    assert info.reason == "changes from back.xlsx"
    diffs = bridge.proposal_diff(project, info.id)
    assert diffs[0].path == "Türliste.csv"
    changed = [cell.after for cell in diffs[0].cell_changes]
    assert "EI90" in changed
