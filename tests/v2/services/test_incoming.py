"""Incoming office files stay as received. Working copies are open formats."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from iterthink.services.bridge_real import RealBridge
from iterthink.services.event_log import read

MARK = "Die Fassade ist verputzt."


def _bridge(tmp_path: Path, project_dir: Path) -> tuple[RealBridge, str]:
    bridge = RealBridge(state_dir=tmp_path / "state", machine="m-test")
    return bridge, bridge.make_project(project_dir).id


def _docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(str(path))


def _xlsx(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Türen"
    sheet["A1"] = "Tür"
    sheet["B1"] = "Größe"
    sheet["A2"] = "T01"
    book.save(str(path))


def _pdf() -> bytes:
    header = b"%PDF-1.4\n"
    body = b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    pages = b"2 0 obj<</Type/Pages/Count 1/Kids[3 0 R]>>endobj\n"
    page = b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]>>endobj\n"
    start = len(header) + len(body) + len(pages) + len(page)

    def entry(offset: int) -> bytes:
        return f"{offset:010d} 00000 n \n".encode()

    xref = (
        b"xref\n0 4\n0000000000 65535 f \n"
        + entry(len(header))
        + entry(len(header) + len(body))
        + entry(len(header) + len(body) + len(pages))
    )
    trailer = f"trailer<</Size 4/Root 1 0 R>>\nstartxref\n{start}\n%%EOF\n".encode()
    return header + body + pages + page + xref + trailer


def test_docx_kept_as_received_and_md_copy_linked(tmp_path: Path, project_dir: Path) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    source = tmp_path / "Brief.docx"
    _docx(source, MARK)
    original = source.read_bytes()
    (project_dir / "Brief.docx").write_bytes(original)
    made = bridge.work_on(project, "Brief.docx", "user:simon")
    assert (project_dir / "Brief.docx").read_bytes() == original
    assert MARK in (project_dir / "Brief.md").read_text(encoding="utf-8")
    assert made.path == "Brief.md"
    log = project_dir / ".iterthink" / "log" / "m-test.jsonl"
    event = next(item for item in read(log) if item.payload.get("relation") == "imported_from")
    assert event.payload["path"] == "Brief.md"
    assert event.payload["source"] == "Brief.docx"
    assert event.payload["source_version"]
    assert event.payload["version"] == made.version


def test_scanned_pdf_goes_through_ocr(
    tmp_path: Path, project_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    (project_dir / "scan.pdf").write_bytes(_pdf())
    monkeypatch.setattr("iterthink.services.incoming.classify_pdf_profile", lambda _src: "text")
    seen: list[Path] = []

    def fake_ocr(src: Path) -> str:
        seen.append(src)
        return "text from the scan"

    monkeypatch.setattr("iterthink.services.ocr_import.pdf_to_markdown_ocr", fake_ocr)
    made = bridge.work_on(project, "scan.pdf", "user:simon")
    assert seen
    assert seen[0].name == "scan.pdf"
    assert made.path == "scan.md"
    assert "text from the scan" in (project_dir / "scan.md").read_text(encoding="utf-8")


def test_xlsx_becomes_csv_working_copy(tmp_path: Path, project_dir: Path) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    source = tmp_path / "Tuerliste.xlsx"
    _xlsx(source)
    original = source.read_bytes()
    (project_dir / "Tuerliste.xlsx").write_bytes(original)
    made = bridge.work_on(project, "Tuerliste.xlsx", "user:simon")
    assert (project_dir / "Tuerliste.xlsx").read_bytes() == original
    text = (project_dir / "Tuerliste.csv").read_text(encoding="utf-8")
    assert "Tür" in text
    assert "Größe" in text
    assert made.path == "Tuerliste.csv"


def test_dropped_file_with_existing_name_opens_review(tmp_path: Path, project_dir: Path) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    bridge.save(project, "Notiz.md", "alt\n", "user:simon", "first")
    dropped = bridge.drop_file(project, "Notiz.md", b"neu\n", "user:simon")
    assert (project_dir / "Notiz.md").read_bytes() == b"neu\n"
    assert "/diff" in dropped.review
    assert "Notiz.md" in dropped.review
    diff = bridge.diff(project, "Notiz.md", None, dropped.version)
    assert any("neu" in block.after for block in diff.blocks)


def test_plan_pdf_is_never_converted(
    tmp_path: Path, project_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    original = _pdf()
    (project_dir / "Grundriss.pdf").write_bytes(original)

    def fail_ocr(_src: Path) -> str:
        raise AssertionError("a plan is not converted")

    monkeypatch.setattr("iterthink.services.ocr_import.pdf_to_markdown_ocr", fail_ocr)
    made = bridge.work_on(project, "Grundriss.pdf", "user:simon")
    assert made.path == ""
    assert made.kind == "plan"
    assert not (project_dir / "Grundriss.md").exists()
    assert (project_dir / "Grundriss.pdf").read_bytes() == original
