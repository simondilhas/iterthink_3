"""Delivery files from Markdown and CSV. They are not versions of the source."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from iterthink.persistence.snapshots import Snapshots
from iterthink.services.bridge_real import RealBridge
from iterthink.services.event_log import read_all
from iterthink.web.app import TOKEN_HEADER, create_app

_CSV = "name,note\nTür,Größe\n"


def _bridge(tmp_path: Path, project_dir: Path) -> tuple[RealBridge, str]:
    bridge = RealBridge(state_dir=tmp_path / "state", machine="m-test")
    return bridge, bridge.make_project(project_dir).id


def _template(folder: Path) -> Path:
    path = folder / "client.docx"
    document = Document()
    document.add_paragraph("CLIENT-MARK")
    document.save(str(path))
    return path


def test_docx_from_markdown_with_client_template(tmp_path: Path, project_dir: Path) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    template = _template(tmp_path)
    info = bridge.export(project, "Pflichtenheft.md", "docx", template=str(template))
    assert info.format == "docx"
    assert info.source_version == Snapshots(project_dir).head().id
    assert info.path.startswith(".iterthink/exports/Pflichtenheft-")
    assert info.path.endswith(".docx")
    written = project_dir / info.path
    assert written.is_file()
    paragraphs = [item.text for item in Document(str(written)).paragraphs]
    assert "CLIENT-MARK" in paragraphs
    assert any("EI30" in item for item in paragraphs)


def test_xlsx_from_csv_keeps_umlauts(tmp_path: Path, project_dir: Path) -> None:
    (project_dir / "liste.csv").write_text(_CSV, encoding="utf-8")
    bridge, project = _bridge(tmp_path, project_dir)
    info = bridge.export(project, "liste.csv", "xlsx")
    assert info.format == "xlsx"
    assert info.path.startswith(".iterthink/exports/liste-")
    book = load_workbook(project_dir / info.path)
    values = [cell.value for row in book.active.iter_rows() for cell in row]
    assert "Tür" in values
    assert "Größe" in values


def test_exports_are_not_versioned_and_linked_to_source_version(
    tmp_path: Path, project_dir: Path
) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    info = bridge.export(project, "Pflichtenheft.md", "pdf")
    assert (project_dir / info.path).read_bytes().startswith(b"%PDF")
    snaps = Snapshots(project_dir)
    assert info.source_version == snaps.head().id
    assert info.path not in snaps.head().paths
    assert all(info.path not in version.paths for version in snaps.history())
    linked = [
        event
        for event in read_all(project_dir)
        if event.type == "export" and event.payload.get("path") == info.path
    ]
    assert len(linked) == 1
    assert linked[0].payload["relation"] == "derived_from"
    assert linked[0].payload["source"] == "Pflichtenheft.md"
    assert linked[0].payload["source_version"] == info.source_version
    (project_dir / "extra.md").write_text("later\n", encoding="utf-8")
    from iterthink.services.projects import tracked_files

    snaps.snapshot(tracked_files(project_dir), "another edit", "user:ada")
    assert info.path not in snaps.head().paths


def test_deliver_page_names_the_file(tmp_path: Path, project_dir: Path) -> None:
    bridge, project = _bridge(tmp_path, project_dir)
    client = TestClient(create_app(bridge=bridge, token="token", actor="user:ada"))
    page = client.get(f"/p/{project}/exports", headers={TOKEN_HEADER: "token"})
    assert page.status_code == 200
    assert "Pflichtenheft.md" in page.text
    posted = client.post(
        f"/p/{project}/exports",
        data={"path": "Pflichtenheft.md", "format": "pdf", "template": ""},
        headers={TOKEN_HEADER: "token"},
    )
    assert posted.status_code == 200
    assert ".iterthink/exports/Pflichtenheft-" in posted.text
    assert "from version" in posted.text
