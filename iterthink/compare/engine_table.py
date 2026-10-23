"""CSV, TSV and XLSX diffs. Daff aligns rows; the first column is the row key."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import daff
from openpyxl import load_workbook

from iterthink.services.bridge import CellChange, FileDiff, FileKind


class TableEngine:
    """One sheet or one table. Cell values are strings."""

    kinds: tuple[FileKind, ...] = ("table",)

    def diff(self, path: str, before: bytes | None, after: bytes | None) -> FileDiff:
        old = _table(path, before)
        new = _table(path, after)
        columns, changes, added, removed, column_added = _align(old, new)
        summary = _summary(changes, added, removed, column_added)
        return FileDiff(
            path,
            "table",
            summary,
            added_rows=tuple(added),
            removed_rows=tuple(removed),
            cell_changes=tuple(changes),
            columns=tuple(columns),
        )


def _table(path: str, data: bytes | None) -> list[list[str]]:
    if not data:
        return []
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return _xlsx(data)
    text = data.decode("utf-8-sig")
    dialect = "excel-tab" if suffix == ".tsv" else "excel"
    return [[cell for cell in row] for row in csv.reader(io.StringIO(text), dialect=dialect)]


def _xlsx(data: bytes) -> list[list[str]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if value is None else str(value) for value in row])
        return rows
    finally:
        workbook.close()


def _align(
    old: list[list[str]], new: list[list[str]]
) -> tuple[list[str], list[CellChange], list[tuple[str, ...]], list[tuple[str, ...]], int]:
    if not old and not new:
        return [], [], [], [], 0
    highlighted: list[list[object]] = []
    alignment = daff.Coopy.compareTables(
        daff.PythonTableView(old), daff.PythonTableView(new)
    ).align()
    daff.TableDiff(alignment, daff.CompareFlags()).hilite(daff.PythonTableView(highlighted))
    return _read_highlight(highlighted)


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _read_highlight(
    table: list[list[object]],
) -> tuple[list[str], list[CellChange], list[tuple[str, ...]], list[tuple[str, ...]], int]:
    header = next((row for row in table if row and row[0] == "@@"), None)
    if header is None:
        return [], [], [], [], 0
    columns = [_cell(value) for value in header[1:]]
    bang = next((row for row in table if row and row[0] == "!"), ["!"])
    actions = [_cell(value) for value in bang[1:]]
    column_added = sum(1 for action in actions if action == "+++")
    changes: list[CellChange] = []
    added: list[tuple[str, ...]] = []
    removed: list[tuple[str, ...]] = []
    for row in table:
        if not row or row[0] in {"!", "@@"}:
            continue
        marker = _cell(row[0])
        cells = [_cell(value) for value in row[1:]]
        cells.extend("" for _ in range(len(columns) - len(cells)))
        if marker == "+++":
            added.append(tuple(cells))
        elif marker == "---":
            removed.append(tuple(cells))
        elif marker in {"->", "+"}:
            changes.extend(_cell_changes(columns, actions, cells))
    return columns, changes, added, removed, column_added


def _cell_changes(columns: list[str], actions: list[str], cells: list[str]) -> list[CellChange]:
    key = cells[0].split("->", 1)[0] if cells else ""
    found: list[CellChange] = []
    for index, column in enumerate(columns):
        raw = cells[index] if index < len(cells) else ""
        action = actions[index] if index < len(actions) else ""
        if "->" in raw:
            before, after = raw.split("->", 1)
            found.append(CellChange(key, column, before, after))
        elif action == "+++":
            found.append(CellChange(key, column, "", raw))
        elif action == "---":
            found.append(CellChange(key, column, raw, ""))
    return found


def _summary(
    changes: list[CellChange],
    added: list[tuple[str, ...]],
    removed: list[tuple[str, ...]],
    column_added: int,
) -> str:
    parts: list[str] = []
    if changes:
        parts.append(f"{len(changes)} cell{'s' if len(changes) != 1 else ''} changed")
    if column_added:
        parts.append(f"{column_added} column{'s' if column_added != 1 else ''} added")
    if added:
        parts.append(f"{len(added)} row{'s' if len(added) != 1 else ''} added")
    if removed:
        parts.append(f"{len(removed)} row{'s' if len(removed) != 1 else ''} removed")
    return ", ".join(parts) if parts else "no changes"
