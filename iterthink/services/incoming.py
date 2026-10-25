"""Keep a received office file, and make an open working copy when asked."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol, cast
from urllib.parse import quote

from iterthink.persistence.snapshots import SnapshotError, Snapshots
from iterthink.services.bridge import Forbidden, NotFound, VersionInfo
from iterthink.services.document_import import (
    classify_pdf_profile,
    docx_to_markdown,
    pdf_to_markdown,
)
from iterthink.services.projects import inside, relative, shown, tracked_files, write_bytes

_EMPTY = re.compile(r"<!--\s*page:\d+\s*-->")


class _Host(Protocol):
    def _root(self, project: str) -> Path: ...
    def _event(
        self, root: Path, project: str, actor: str, kind: str, payload: dict[str, object]
    ) -> None: ...
    def save(
        self, project: str, path: str, text: str, actor: str, message: str = ""
    ) -> VersionInfo | None: ...


@dataclass(frozen=True)
class Worked:
    """A working copy, or a plan that stays as the PDF."""

    path: str
    version: str
    kind: str


@dataclass(frozen=True)
class Dropped:
    """Where the dropped file landed. ``placed`` is false until a folder is chosen."""

    review: str
    version: str
    placed: bool


def _host(bridge: IncomingBridge) -> _Host:
    return cast(_Host, bridge)


class IncomingBridge:
    """Real-bridge import of a file someone else sent."""

    def work_on(self, project: str, path: str, actor: str) -> Worked:
        host = _host(self)
        root = host._root(project)
        rel, full = _file(root, path)
        source = _ensure(host, root, project, rel, full.read_bytes(), actor)
        if full.suffix.lower() == ".pdf" and classify_pdf_profile(full) == "plan":
            return Worked("", source, "plan")
        text, dest = _working_copy(full, rel)
        saved = host.save(project, dest, text, actor, f"working copy of {rel}")
        version = saved.id if saved is not None else Snapshots(root).head().id
        host._event(root, project, actor, "relation", _link(dest, version, rel, source))
        kind = "table" if dest.endswith(".csv") else "markdown"
        return Worked(dest, version, kind)

    def drop_file(
        self, project: str, name: str, data: bytes, actor: str, folder: str | None = None
    ) -> Dropped:
        host = _host(self)
        root = host._root(project)
        rel = _destination(root, name, folder)
        if rel is None:
            return Dropped("", "", False)
        version = _ensure(host, root, project, rel, data, actor)
        return Dropped(_review(project, rel, version), version, True)


def _file(root: Path, path: str) -> tuple[str, Path]:
    rel = relative(path)
    full = inside(root, rel)
    if not full.is_file():
        raise NotFound(f"no file {rel}")
    return rel, full


def _working_copy(full: Path, rel: str) -> tuple[str, str]:
    suffix = full.suffix.lower()
    dest = str(Path(rel).with_suffix(".md" if suffix != ".xlsx" else ".csv"))
    if suffix == ".docx":
        assets = full.parent / f"{full.stem}_assets"
        return docx_to_markdown(full, assets), dest
    if suffix == ".pdf":
        return _pdf_text(full), dest
    if suffix == ".xlsx":
        return _sheet_csv(full), dest
    raise Forbidden(f"{rel} is already an open file")


def _pdf_text(full: Path) -> str:
    try:
        text = pdf_to_markdown(full)
    except (OSError, ValueError, RuntimeError):
        text = ""
    if _has_body(text):
        return text
    from iterthink.services.ocr_import import pdf_to_markdown_ocr

    return pdf_to_markdown_ocr(full)


def _has_body(text: str) -> bool:
    if "PDF text extraction returned no content" in text:
        return False
    return bool(_EMPTY.sub("", text).strip())


def _sheet_csv(full: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(full, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\n")
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if cell is None else cell for cell in row])
        return buffer.getvalue()
    finally:
        book.close()


def _link(path: str, version: str, source: str, source_version: str) -> dict[str, object]:
    return {
        "path": path,
        "relation": "imported_from",
        "source": source,
        "source_version": source_version,
        "version": version,
    }


def _ensure(host: _Host, root: Path, project: str, rel: str, data: bytes, actor: str) -> str:
    full = inside(root, rel)
    write_bytes(full, data)
    snaps = Snapshots(root)
    current = _current(snaps, rel, data)
    if current is not None:
        return current
    made = snaps.snapshot(tracked_files(root), f"received {rel}", actor)
    version = shown(made if made is not None else snaps.head())
    host._event(
        root,
        project,
        actor,
        "version",
        {"kind": "received", "path": rel, "version": version.id},
    )
    return version.id


def _current(snaps: Snapshots, rel: str, data: bytes) -> str | None:
    try:
        head = snaps.head()
    except SnapshotError:
        return None
    if snaps.file_at(head.id, rel) == data:
        return head.id
    return None


def _destination(root: Path, name: str, folder: str | None) -> str | None:
    rel = relative(Path(name).name if folder is not None else name)
    if inside(root, rel).is_file() and folder is None:
        return rel
    if folder is None:
        return None
    prefix = "" if folder in {"", "."} else relative(folder)
    return f"{prefix}/{rel}" if prefix else rel


def _review(project: str, rel: str, version: str) -> str:
    return f"/p/{project}/diff?path={quote(rel)}&target={version}"
