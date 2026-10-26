"""An edited Word or Excel file comes back as one proposal to the open source."""

from __future__ import annotations

import csv
import io
from collections.abc import Mapping
from pathlib import Path
from typing import Protocol, cast

from docx import Document
from openpyxl import load_workbook

from iterthink.compare.engine_text import TextEngine
from iterthink.persistence.snapshots import Snapshots
from iterthink.services.bridge import Conflict, Invalid, NotFound, ProposalInfo
from iterthink.services.projects import inside, relative, tracked_files


class _Caller(Protocol):
    def _root(self, project: str) -> Path: ...
    def _event(
        self, root: Path, project: str, actor: str, kind: str, payload: dict[str, object]
    ) -> None: ...
    def propose(
        self,
        project: str,
        changes: Mapping[str, str],
        actor: str,
        reason: str,
        baseline: str | None = None,
    ) -> ProposalInfo: ...


def take_back(host: object, project: str, returned: Path, source: str, actor: str) -> ProposalInfo:
    """Keep the received file and propose its text to the Markdown or CSV source."""
    caller = cast(_Caller, host)
    root = caller._root(project)
    rel = relative(source)
    full = inside(root, rel)
    if not full.is_file():
        raise NotFound(f"no file {rel}")
    if not returned.is_file():
        raise NotFound(f"no file {returned}")
    kept = _keep(root, returned)
    _version_received(caller, root, project, kept, actor)
    current = full.read_text(encoding="utf-8")
    proposed = _proposed(rel, current, kept)
    if _same(proposed, current):
        raise Invalid("the returned file matches the source")
    reason = f"changes from {kept.name}"
    return caller.propose(project, {rel: proposed}, actor, reason)


def _keep(root: Path, returned: Path) -> Path:
    resolved = returned.resolve()
    base = root.resolve()
    if resolved.is_relative_to(base):
        rel = resolved.relative_to(base)
        if any(part.startswith(".") for part in rel.parts):
            raise Invalid("a returned file stays a source file")
        return resolved
    dest = base / returned.name
    incoming = returned.read_bytes()
    if dest.is_file() and dest.read_bytes() != incoming:
        raise Conflict(f"{returned.name} is already in the project")
    if not dest.exists():
        dest.write_bytes(incoming)
    return dest


def _version_received(host: _Caller, root: Path, project: str, kept: Path, actor: str) -> None:
    rel = kept.resolve().relative_to(root.resolve()).as_posix()
    snaps = Snapshots(root)
    if rel in snaps.head().paths:
        return
    made = snaps.snapshot(tracked_files(root), f"received {kept.name}", actor)
    if made is None:
        return
    host._event(
        root, project, actor, "version", {"kind": "received", "path": rel, "version": made.id}
    )


def _proposed(rel: str, current: str, kept: Path) -> str:
    suffix = kept.suffix.lower()
    if suffix == ".docx" and Path(rel).suffix.lower() in {".md", ".markdown", ".txt"}:
        return _from_text(rel, current, _docx_text(kept))
    if suffix == ".xlsx" and Path(rel).suffix.lower() == ".csv":
        return _csv_text(kept)
    raise Invalid("a returned file is Word for Markdown, or Excel for CSV")


def _from_text(rel: str, current: str, extracted: str) -> str:
    diff = TextEngine().diff(rel, current.encode("utf-8"), extracted.encode("utf-8"))
    parts = [block.after for block in diff.blocks if block.after]
    text = "\n\n".join(parts)
    if text and not text.endswith("\n"):
        text += "\n"
    return text


def _docx_text(path: Path) -> str:
    parts = [item.text.strip() for item in Document(str(path)).paragraphs if item.text.strip()]
    return "\n\n".join(parts)


def _csv_text(path: Path) -> str:
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        rows = [
            ["" if value is None else str(value) for value in row]
            for row in book.worksheets[0].iter_rows(values_only=True)
        ]
    finally:
        book.close()
    buffer = io.StringIO(newline="")
    csv.writer(buffer, lineterminator="\n").writerows(rows)
    return buffer.getvalue()


def _same(proposed: str, current: str) -> bool:
    return proposed.replace("\r\n", "\n").strip() == current.replace("\r\n", "\n").strip()
